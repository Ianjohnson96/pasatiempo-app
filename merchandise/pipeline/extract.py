"""Turn POS reports (PDF, Excel or CSV) into text.

The Pro Shop POS prints most reports portrait (pdfplumber reads them cleanly) but the
SKU Analysis and Rounds Summary are printed rotated; pdfplumber returns reversed
fragments for those, so we fall back to pypdf, which honours the text matrix.

A spreadsheet export becomes one line per row with its cells separated by tabs, the
numbers written the way the printout shows them, so the same parsers read both.
"""
import csv
from datetime import date, datetime
import io

import pdfplumber
from pypdf import PdfReader

MARKERS = ('Pasatiempo', 'Net Sales', 'Daily Sales')


def pdf_text(src):
    """src: a file path, or the file's bytes."""
    source = (lambda: io.BytesIO(src)) if isinstance(src, (bytes, bytearray)) else (lambda: src)
    with pdfplumber.open(source()) as pdf:
        # Decide from the first page, so a rotated report isn't read twice (the SKU Analysis
        # takes far longer through pdfplumber than through pypdf).
        first = (pdf.pages[0].extract_text() or '') if pdf.pages else ''
        if len(first) >= 500 and not any(m in first[:500] for m in MARKERS):
            text = None
        else:
            text = '\n'.join([first] + [(p.extract_text() or '') for p in pdf.pages[1:]])
            if not any(m in text[:500] for m in MARKERS):
                text = None
    if text is not None:
        return text
    return '\n'.join(p.extract_text() or '' for p in PdfReader(source()).pages)


def detect_kind(text):
    head = text[:1500]
    if 'SKU Analysis as of' in head:
        return 'sku_analysis'
    if 'BEST' in head and 'Quantity Sold' in head:
        return 'best100'
    if 'Net Sales by Category' in head:
        return 'sales_by_category'
    if 'Net Sales by Item' in head:
        return 'sales_by_item'
    if 'Daily Sales Report' in head:
        return 'daily_sales'
    if 'Rounds Summary' in head:
        return 'rounds'
    return None


SHEETS = ('.xlsx', '.xlsm', '.xls', '.csv')


def is_sheet(name):
    return name.lower().endswith(SHEETS)


def _cell(v, fmt=''):
    """One spreadsheet cell as the printed report shows it."""
    if v is None:
        return ''
    if isinstance(v, (datetime, date)):
        return f'{v:%b}{v.day:02d}/{v:%y}'  # Sep28/18, as the POS prints dates
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        if '%' in (fmt or ''):
            return f'{v * 100:.1f}%'
        sep = ',' if ',' in (fmt or '') else ''  # no separator unless the cell shows one (years stay 2026)
        if '.0' in (fmt or '') or (isinstance(v, float) and v != int(v)):
            return f'{v:{sep}.2f}'
        return f'{int(v):{sep}}'
    return ' '.join(str(v).split())


def _lines(rows):
    out = []
    for cells in rows:
        cells = list(cells)
        while cells and cells[-1] == '':
            cells.pop()
        while cells and cells[0] == '':
            cells.pop(0)
        # an empty cell between values stands where the printout has a dash
        out.append('\t'.join(c if c else '-' for c in cells))
    return '\n'.join(out)


def sheet_text(name, data):
    """name: the file name (for its type); data: the file's bytes."""
    n = name.lower()
    if n.endswith('.csv'):
        try:
            txt = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            txt = data.decode('latin-1')
        return _lines([_cell(c) for c in r] for r in csv.reader(io.StringIO(txt)))
    if n.endswith('.xls'):
        import xlrd
        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
        rows = []
        for sh in book.sheets():
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append('')
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(_cell(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
                    else:
                        fmt = book.format_map.get(book.xf_list[cell.xf_index].format_key)
                        row.append(_cell(cell.value, fmt.format_str if fmt else ''))
                rows.append(row)
        return _lines(rows)
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = []
    for ws in book.worksheets:
        for r in ws.iter_rows():
            rows.append([_cell(c.value, getattr(c, 'number_format', '')) for c in r])
    book.close()
    return _lines(rows)


def report_text(name, src):
    """Text of one report file. src: a path, or the file's bytes."""
    if is_sheet(name):
        if not isinstance(src, (bytes, bytearray)):
            with open(src, 'rb') as f:
                src = f.read()
        return sheet_text(name, src)
    return pdf_text(src)
